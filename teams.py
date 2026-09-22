"""주간 이슈를 Microsoft Teams 채널에 게시한다.

매주 하는 일
  1) 주간 Word 보고서 생성: docs/reports/<주차>.docx (웹에서도 내려받을 수 있게 docs 아래 둔다)
  2) 채널 파일 폴더에 업로드: OneDrive로 동기화한 SharePoint 폴더(config teams.sync_dir)에 복사
     → OneDrive 앱이 채널 '파일' 탭으로 올린다. 계정 비밀번호·API 키를 코드가 다루지 않는다.
  3) 채널 게시물(게시판): Teams '워크플로' 웹훅(.env TEAMS_WEBHOOK_URL)으로 카드 게시

설정된 것만 수행하며, 설정된 단계가 모두 성공해야 그 주차를 '완료'로 기록한다.
"""
from __future__ import annotations

import json
import os
import shutil
import urllib.request
from datetime import datetime, timedelta

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from curate import FIELD_ORDER

BASE = os.path.dirname(os.path.abspath(__file__))
REPORT_DIR = os.path.join(BASE, "docs", "reports")
FONT = "맑은 고딕"
ACCENT = RGBColor(0x5B, 0x4B, 0xFF)
MUTED = RGBColor(0x6B, 0x6B, 0x78)


# ---------------------------------------------------------------- Word 보고서
def _font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def _link(paragraph, url: str, text: str):
    """문단에 하이퍼링크를 넣는다 (python-docx에 공식 API가 없어 oxml로 구성)."""
    rid = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "5B4BFF")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    for el in (fonts, color, u):
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def _heading(doc, text, level):
    h = doc.add_heading(level=level)
    _font(h.add_run(text), size={0: 20, 1: 15, 2: 12.5}[level], bold=True,
          color=ACCENT if level < 2 else None)
    return h


def _labeled(doc, label, text):
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _font(p.add_run(f"{label}  "), size=10, bold=True, color=ACCENT)
    _font(p.add_run(" ".join(str(text).split())), size=10)


def _field_order(papers):
    groups: dict[str, list[dict]] = {}
    for p in papers:
        groups.setdefault(p.get("subfield", "기타 AI"), []).append(p)
    order = [f for f in FIELD_ORDER if f in groups] + [f for f in groups if f not in FIELD_ORDER]
    return [(f, sorted(groups[f], key=lambda x: x.get("upvotes", 0), reverse=True)) for f in order]


def _window(issue: dict) -> list[str]:
    """이슈 대상 기간. 기간 기록이 없는 과거 이슈는 발행일 직전 7일로 추정한다."""
    if issue.get("window"):
        return issue["window"]
    try:
        d = datetime.fromisoformat(issue.get("date", "").replace("Z", "+00:00")).date()
        return [(d - timedelta(days=7)).isoformat(), (d - timedelta(days=1)).isoformat()]
    except Exception:  # noqa: BLE001
        return ["", ""]


def _authors(p: dict, n: int = 6) -> str:
    names = [a for a in p.get("authors", []) if any(ch.isalnum() for ch in a)]
    return ", ".join(names[:n]) + (" 외" if len(names) > n else "")


def build_report(label: str, issue: dict, papers: list[dict], cfg: dict) -> str:
    os.makedirs(REPORT_DIR, exist_ok=True)
    site = cfg.get("site_url", "").rstrip("/")
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Pt(56)

    _heading(doc, f"AI 논문 요약 다이제스트 · {label}", 0)
    win = _window(issue)
    meta = doc.add_paragraph()
    _font(meta.add_run(
        f"대상 기간 {win[0]} ~ {win[1]}  ·  {len(papers)}편  ·  "
        f"{len({p.get('subfield') for p in papers})}개 분야  ·  "
        f"작성 {datetime.now().strftime('%Y-%m-%d')}"), size=9.5, color=MUTED)
    if site:
        p = doc.add_paragraph()
        _font(p.add_run("웹에서 보기: "), size=9.5, color=MUTED)
        _link(p, f"{site}/issue/{label}.html", f"{site}/issue/{label}.html")
    note = doc.add_paragraph()
    _font(note.add_run(
        "선별 기준: 대상 기간에 Hugging Face Daily Papers에 오른 논문을 AI 세부 분야별로 묶어 "
        "커뮤니티 추천수(👍) 상위 2~3편씩 골랐습니다. 요약은 AI가 초록을 바탕으로 작성했으므로 "
        "중요한 판단은 원문으로 확인하세요."), size=9, color=MUTED)

    top = sorted(papers, key=lambda x: x.get("upvotes", 0), reverse=True)[:3]
    _heading(doc, "이번 주 인기 1~3위", 1)
    for i, p in enumerate(top, 1):
        s = p.get("summary", {})
        _heading(doc, f"{i}위 · {p['title']}  (👍 {p.get('upvotes', 0)})", 2)
        _labeled(doc, "주제", s.get("tldr"))
        _labeled(doc, "핵심 원리 및 기술", s.get("method_easy") or s.get("method"))
        _labeled(doc, "개선효과", s.get("improvement") or s.get("result"))
        _labeled(doc, "미비점 및 보완방안", s.get("limitations"))

    for field, items in _field_order(papers):
        _heading(doc, f"{field}  ({len(items)}편)", 1)
        for p in items:
            s = p.get("summary", {})
            _heading(doc, p["title"], 2)
            m = doc.add_paragraph()
            _font(m.add_run(f"{_authors(p)}  ·  👍 {p.get('upvotes', 0)}  ·  게재 {str(p.get('published', ''))[:10]}  ·  "),
                  size=9, color=MUTED)
            _link(m, p.get("arxiv_url", ""), f"arXiv:{p['id']}")
            _labeled(doc, "한 줄 요약", s.get("tldr"))
            _labeled(doc, "12살 버전", s.get("eli12"))
            _labeled(doc, "예시", s.get("example"))
            _labeled(doc, "어떤 문제를 푸는가", s.get("problem"))
            _labeled(doc, "핵심 원리 및 기술", s.get("method_easy") or s.get("method"))
            _labeled(doc, "개선효과", s.get("improvement") or s.get("result"))
            _labeled(doc, "미비점 및 보완방안", s.get("limitations"))
            if s.get("background"):
                b = doc.add_paragraph()
                _font(b.add_run("배경지식"), size=10, bold=True, color=ACCENT)
                for item in s["background"]:
                    li = doc.add_paragraph(style="List Bullet")
                    _font(li.add_run(f"{item.get('term', '')}: "), size=9.5, bold=True)
                    _font(li.add_run(item.get("explain", "")), size=9.5)
            if s.get("keywords"):
                _labeled(doc, "키워드", ", ".join(s["keywords"]))

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(foot.add_run("AI 논문 요약 다이제스트 · 자동 생성"), size=8.5, color=MUTED)

    path = os.path.join(REPORT_DIR, f"{label}.docx")
    tmp = path + ".tmp"
    doc.save(tmp)
    os.replace(tmp, path)
    return path


# ---------------------------------------------------------------- 채널 파일 업로드
def upload_to_folder(report: str, label: str, tcfg: dict) -> bool | None:
    """OneDrive 동기화 폴더에 복사. 미설정이면 None."""
    sync_dir = os.path.expanduser(tcfg.get("sync_dir", "") or "")
    if not sync_dir:
        return None
    if not os.path.isdir(sync_dir):
        print(f"  Teams 폴더 없음(OneDrive 동기화 확인 필요): {sync_dir}", flush=True)
        return False
    dest_dir = os.path.join(sync_dir, tcfg.get("subfolder", "AI 논문 다이제스트"))
    try:
        os.makedirs(dest_dir, exist_ok=True)
        dest = os.path.join(dest_dir, f"{label}_AI논문다이제스트.docx")
        shutil.copy2(report, dest)
        print(f"  Teams 채널 폴더에 업로드: {dest}", flush=True)
        return True
    except Exception as e:  # noqa: BLE001
        print(f"  Teams 폴더 복사 실패: {e}", flush=True)
        return False


# ---------------------------------------------------------------- 채널 게시물
def _text(text, **kw):
    return {"type": "TextBlock", "text": text, "wrap": True, **kw}


def build_card(label: str, issue: dict, papers: list[dict], cfg: dict) -> dict:
    site = cfg.get("site_url", "").rstrip("/")
    tcfg = cfg.get("teams", {})
    win = _window(issue)
    body = [
        _text(f"📚 AI 논문 요약 다이제스트 · {label}", size="Large", weight="Bolder"),
        _text(f"대상 기간 {win[0]} ~ {win[1]} · {len(papers)}편 · "
              f"{len({p.get('subfield') for p in papers})}개 분야", isSubtle=True, spacing="None"),
        _text("🏆 이번 주 인기 1~3위", weight="Bolder", spacing="Medium"),
    ]
    medals = ["🥇", "🥈", "🥉"]
    for i, p in enumerate(sorted(papers, key=lambda x: x.get("upvotes", 0), reverse=True)[:3]):
        s = p.get("summary", {})
        items = [_text(f"{medals[i]} {i + 1}위 · 👍 {p.get('upvotes', 0)}", weight="Bolder"),
                 _text(p["title"], isSubtle=True, spacing="None", size="Small")]
        for k, v in (("주제", s.get("tldr")),
                     ("핵심 원리 및 기술", s.get("method_easy") or s.get("method")),
                     ("개선효과", s.get("improvement") or s.get("result")),
                     ("미비점 및 보완방안", s.get("limitations"))):
            if v:
                items.append(_text(f"**{k}**  {' '.join(str(v).split())}", spacing="Small"))
        body.append({"type": "Container", "style": "emphasis", "spacing": "Medium", "items": items})

    lines = [f"**{f}** · " + " / ".join(x["title"][:48] for x in items) for f, items in _field_order(papers)]
    body.append(_text("분야별 선정 논문", weight="Bolder", spacing="Medium"))
    body.append(_text("\n\n".join(lines), size="Small"))

    actions = []
    if site:
        actions.append({"type": "Action.OpenUrl", "title": "웹에서 보기", "url": f"{site}/issue/{label}.html"})
        actions.append({"type": "Action.OpenUrl", "title": "주간 보고서(Word)", "url": f"{site}/reports/{label}.docx"})
    if tcfg.get("folder_url"):
        actions.append({"type": "Action.OpenUrl", "title": "채널 폴더", "url": tcfg["folder_url"]})
    return {
        "type": "message",
        "attachments": [{
            "contentType": "application/vnd.microsoft.card.adaptive",
            "contentUrl": None,
            "content": {
                "$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
                "type": "AdaptiveCard", "version": "1.4", "msteams": {"width": "Full"},
                "body": body, "actions": actions,
            },
        }],
    }


def post_card(card: dict) -> bool | None:
    """Teams 워크플로 웹훅으로 카드 게시. 미설정이면 None."""
    url = os.getenv("TEAMS_WEBHOOK_URL")
    if not url:
        return None
    req = urllib.request.Request(url, data=json.dumps(card, ensure_ascii=False).encode("utf-8"),
                                 headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            ok = 200 <= resp.status < 300
        print(f"  Teams 채널 게시 {'완료' if ok else f'실패(HTTP {resp.status})'}", flush=True)
        return ok
    except Exception as e:  # noqa: BLE001
        print(f"  Teams 채널 게시 실패: {e}", flush=True)
        return False


# ---------------------------------------------------------------- 진입점
def report_path(label: str) -> str:
    return os.path.join(REPORT_DIR, f"{label}.docx")


def configured(cfg: dict) -> set[str]:
    """지금 설정돼 있는 Teams 전달 단계."""
    steps = set()
    if (cfg.get("teams", {}).get("sync_dir") or "").strip():
        steps.add("folder")
    if os.getenv("TEAMS_WEBHOOK_URL"):
        steps.add("post")
    return steps


def pending(issue: dict, cfg: dict) -> set[str]:
    """설정돼 있지만 이 주차에서 아직 성공하지 못한 단계.

    단계별로 따로 기록하므로, 나중에 새 단계를 설정해도(예: OneDrive 동기화 추가)
    지난 주차까지 그 단계만 이어서 수행된다.
    """
    done = issue.get("teams_steps", {})
    return {s for s in configured(cfg) if not done.get(s)}


def publish(label: str, issue: dict, papers: list[dict], cfg: dict) -> bool:
    """남은 Teams 단계(폴더 업로드, 채널 게시)를 수행. 남은 단계가 없으면 True."""
    todo = pending(issue, cfg)
    if not papers or not todo:
        return not todo
    report = report_path(label)
    if not os.path.exists(report):
        build_report(label, issue, papers, cfg)
    steps = issue.setdefault("teams_steps", {})
    if "folder" in todo:
        steps["folder"] = bool(upload_to_folder(report, label, cfg.get("teams", {})))
    if "post" in todo:
        steps["post"] = bool(post_card(build_card(label, issue, papers, cfg)))
    return not pending(issue, cfg)
